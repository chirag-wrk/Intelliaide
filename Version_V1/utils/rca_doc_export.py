"""Build RCA stage exports as Word HTML (.doc) and OOXML (.docx) with causal-DAG images."""

from __future__ import annotations

import html
import io
import re

from docx import Document
from docx.shared import Inches, Pt

# Letter page content area with 1 in margins.
_PAGE_CONTENT_WIDTH_IN = 6.5
_PAGE_CONTENT_MAX_HEIGHT_IN = 8.0


def split_after_heading(text: str, heading_title: str) -> tuple[str, str]:
    """Split markdown after a ## section (matches frontend splitAfterHeading)."""
    if not text:
        return text, ""
    lines = text.split("\n")
    target = heading_title.strip().lower()
    found_target = False
    split_index = -1
    for i, line in enumerate(lines):
        m = re.match(r"^##\s+(.+)", line)
        if not m:
            continue
        title = m.group(1).strip().lower()
        if found_target:
            split_index = i
            break
        if title == target:
            found_target = True
    if not found_target or split_index == -1:
        return text, ""
    return "\n".join(lines[:split_index]).rstrip(), "\n".join(lines[split_index:]).lstrip()


def _text_block_to_html(text: str) -> str:
    return html.escape(text or "").replace("\n", "<br/>")


def _causal_dag_embed_html(image_rel_path: str) -> str:
    """Word/Google-Docs-safe HTML image reference via a sibling *_files folder."""
    safe_src = html.escape(image_rel_path, quote=True)
    return (
        '<h2 style="font-size:13pt;color:#1e293b;margin:24px 0 8px 0;'
        'border-bottom:1px solid #e2e8f0;padding-bottom:4px;">Causal Chain</h2>'
        '<p style="font-size:9pt;color:#64748b;margin:0 0 12px 0;">'
        "Visual map of how the identified causes led to the reported symptom(s)."
        "</p>"
        f'<img src="{safe_src}" alt="Causal Chain diagram" '
        'style="display:block;width:100%;max-width:6.5in;max-height:8in;'
        'height:auto;object-fit:contain;margin:0 auto 24px auto;" />'
    )


def build_rca_stage_doc_html(
    title: str,
    body_text: str,
    *,
    image_rel_path: str | None = None,
) -> str:
    """Word-compatible HTML saved as .doc, with optional relative-path DAG image."""
    safe_title = html.escape(title)
    if image_rel_path:
        before, after = split_after_heading(body_text, "Executive Summary")
        body_parts = [_text_block_to_html(before), _causal_dag_embed_html(image_rel_path)]
        if after:
            body_parts.append(f"<div>{_text_block_to_html(after)}</div>")
        safe_body = "\n".join(body_parts)
    else:
        safe_body = f"<div>{_text_block_to_html(body_text)}</div>"
    return f"""<!DOCTYPE html>
<html xmlns:o="urn:schemas-microsoft-com:office:office"
xmlns:w="urn:schemas-microsoft-com:office:word"
xmlns="http://www.w3.org/TR/REC-html40">
<head>
  <meta charset="utf-8" />
  <title>{safe_title}</title>
  <style>
    body {{
      font-family: Arial, sans-serif;
      font-size: 11pt;
      line-height: 1.45;
      margin: 1in;
      color: #1f2937;
      white-space: normal;
      word-break: break-word;
    }}
    h1 {{
      font-size: 18pt;
      margin: 0 0 14px 0;
      color: #111827;
    }}
    img {{
      max-width: 6.5in;
      max-height: 8in;
      height: auto;
    }}
  </style>
</head>
<body>
  <h1>{safe_title}</h1>
  <div>{safe_body}</div>
</body>
</html>"""


def _add_markdownish_paragraphs(doc: Document, text: str) -> None:
    if not text:
        return
    for line in text.split("\n"):
        heading2 = re.match(r"^##\s+(.+)", line.strip())
        heading1 = re.match(r"^#\s+(.+)", line.strip())
        if heading2:
            doc.add_heading(heading2.group(1).strip(), level=2)
        elif heading1:
            doc.add_heading(heading1.group(1).strip(), level=1)
        elif line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")


def _add_causal_dag_to_docx(doc: Document, dag_png: bytes) -> None:
    doc.add_heading("Causal Chain", level=2)
    note = doc.add_paragraph(
        "Visual map of how the identified causes led to the reported symptom(s)."
    )
    if note.runs:
        note.runs[0].font.size = Pt(9)
    stream = io.BytesIO(dag_png)
    doc.add_picture(stream, width=Inches(_PAGE_CONTENT_WIDTH_IN))


def build_rca_stage_docx(
    title: str,
    body_text: str,
    dag_png: bytes | None = None,
) -> bytes:
    """OOXML .docx with embedded causal-DAG PNG (works in Word, Google Docs, LibreOffice)."""
    doc = Document()
    doc.add_heading(title, level=1)

    if dag_png:
        before, after = split_after_heading(body_text, "Executive Summary")
        _add_markdownish_paragraphs(doc, before)
        _add_causal_dag_to_docx(doc, dag_png)
        _add_markdownish_paragraphs(doc, after)
    else:
        _add_markdownish_paragraphs(doc, body_text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def append_rca_stage_bundle_entries(
    entries: list[tuple[str, str | bytes]],
    *,
    arc_prefix: str,
    doc_filename: str,
    title: str,
    body_text: str,
    dag_png: bytes | None,
) -> None:
    """Add .doc, .docx, and optional DAG image folder entries for one RCA stage."""
    base = doc_filename.rsplit(".", 1)[0]
    image_rel: str | None = None
    if dag_png:
        image_rel = f"{base}.doc_files/causal_dag.png"
        entries.append((f"{arc_prefix}{image_rel}", dag_png))

    entries.append((
        f"{arc_prefix}{base}.doc",
        "\ufeff" + build_rca_stage_doc_html(title, body_text, image_rel_path=image_rel),
    ))
    entries.append((
        f"{arc_prefix}{base}.docx",
        build_rca_stage_docx(title, body_text, dag_png),
    ))
